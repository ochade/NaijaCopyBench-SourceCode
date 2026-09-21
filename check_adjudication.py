from pathlib import Path
from openpyxl import load_workbook
from collections import Counter

p = Path.home() / "Downloads" / "catA_adjudication_sheet.xlsx"
ws = load_workbook(p, data_only=True).active
hdr = [c.value for c in ws[1]]
i = {h: n for n, h in enumerate(hdr) if h}

PROV = "FINAL_provision_ Sammy"
ANS  = "FINAL_answer_Sammy"
REAS = "reasoning_Sammy"
VERD = "SOUND"

rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
print(f"rows: {len(rows)}\n")

for name, col in [("final provision", PROV), ("final answer", ANS),
                  ("reasoning", REAS), ("verdict", VERD)]:
    n = sum(1 for r in rows if r[i[col]] not in (None, ""))
    print(f"{name:18} filled {n}/{len(rows)}")

print("\nverdict values:")
for k, v in Counter(str(r[i[VERD]]).strip() for r in rows if r[i[VERD]]).most_common():
    print(f"   {k[:40]:42} {v}")

print("\nfirst five:")
for r in rows[:5]:
    print(f"  {r[0]:12} {str(r[i[PROV]])[:26]:28} | {str(r[i[VERD]])[:18]}")
