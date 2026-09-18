import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

rows = [json.loads(l) for l in
        Path("results/m4_verification_sheet.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]

wb = Workbook(); ws = wb.active; ws.title = "M4_verification"
ws.append(["row_id","arm","item_id","category","question","gold_provision",
           "gold_answer","claim","verdict (S/N/C)"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F3864")
    c.alignment = Alignment(wrap_text=True, vertical="center")

rid = 0
for r in rows:
    for ci, claim in enumerate(r["claims"]):
        rid += 1
        first = ci == 0
        ws.append([rid, r["arm"], r["item_id"], r["category"],
                   r["question"][:120] if first else "",
                   r["gold_provision"] if first else "",
                   r["gold_answer"][:300] if first else "",
                   claim, ""])

widths = {"A":7,"B":6,"C":13,"D":14,"E":40,"F":16,"G":46,"H":60,"I":14}
for col,w in widths.items(): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "H2"

out = Path.home()/"Downloads"/"m4_verification.xlsx"
wb.save(out)
print("wrote", out)
print(f"{rid} claims across {len(rows)} responses")
