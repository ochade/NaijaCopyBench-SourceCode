from pathlib import Path
from openpyxl import load_workbook

DL = Path.home()/"Downloads"
FILES = ["catA_adjudication_sheet.xlsx",
         "catB_adjudication_sheet (1).xlsx",
         "catC_review_sheet.xlsx",
         "catD_annotation_sheet_Euchay.xlsx",
         "catE_review_sheet.xlsx",
         "catF_annotation_sheet_Euchay.xlsx"]

for fn in FILES:
    p = DL / fn
    if not p.exists():
        print(f"MISSING: {fn}\n"); continue
    ws = load_workbook(p, data_only=True).active
    hdr = [c.value for c in ws[1]]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    print(f"=== {fn} ===")
    print(f"  rows: {len(rows)}")
    print(f"  headers: {hdr}")
    print(f"  first row: {[str(v)[:30] if v else None for v in rows[0]]}" if rows else "  (empty)")
    print()
