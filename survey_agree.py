from pathlib import Path
from openpyxl import load_workbook
from collections import Counter

DL = Path.home()/"Downloads"

def survey(fn, prov_col, agree_col=None):
    ws = load_workbook(DL/fn, data_only=True).active
    hdr = [c.value for c in ws[1]]
    i = {h: n for n, h in enumerate(hdr) if h}
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    print(f"=== {fn} ({len(rows)} rows) ===")
    filled = sum(1 for r in rows if r[i[prov_col]] not in (None, ""))
    print(f"  {prov_col}: filled {filled}/{len(rows)}")
    if agree_col and agree_col in i:
        vals = Counter(str(r[i[agree_col]]).strip().upper() for r in rows if r[i[agree_col]])
        print(f"  {agree_col} values: {dict(vals)}")
        blank_agreed = sum(1 for r in rows
                           if r[i[prov_col]] in (None, "")
                           and r[i[agree_col]] and "AGREE" in str(r[i[agree_col]]).upper())
        print(f"  blank FINAL where AGREE: {blank_agreed}")
        blank_noagree = sum(1 for r in rows
                            if r[i[prov_col]] in (None, "")
                            and not (r[i[agree_col]] and "AGREE" in str(r[i[agree_col]]).upper()))
        print(f"  blank FINAL, no AGREE either: {blank_noagree}  <-- these have no gold")
    print("  first 5 item_ids:", [r[0] for r in rows[:5]])
    print()

survey("catB_adjudication_sheet (1).xlsx", "FINAL_provision", "basis_of_resolution")
survey("catC_review_sheet.xlsx", "FINAL_provision", "AGREE?")
survey("catE_review_sheet.xlsx", "FINAL_provision", "AGREE?")
