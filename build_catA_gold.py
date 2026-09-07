import re, json
from pathlib import Path
from openpyxl import load_workbook
from collections import Counter

ws = load_workbook(Path.home()/"Downloads"/"catA_adjudication_sheet.xlsx", data_only=True).active
hdr = [c.value for c in ws[1]]
i = {h: n for n, h in enumerate(hdr) if h}

PROV, VERD, REAS = "FINAL_provision_ Sammy", "FINAL_answer_Sammy", "reasoning_Sammy"

# A-011: adjudicator gave s.30, but the question asks about registration's
# evidential effect. Both annotators cited s.87. Overridden.
OVERRIDE = {
 "NCB-A-011": {"provision": "87(4)", "answer_from": "euchay",
   "note": "Adjudicator gave s.30 (assignment). The question concerns registration's "
           "evidential effect, which is s.87(4); both annotators cited s.87. "
           "Researcher override."},
 "NCB-A-002": {"provision": "3(a)", "answer_from": "euchay",
   "note": "Adjudicator's provision accepted. Answer taken from the annotator, whose "
           "statement of the exclusion is fuller than the adjudicator's summary."},
 "NCB-A-015": {"provision": "28(3)", "answer_from": "euchay",
   "note": "Adjudicator's provision accepted. Her reasoning invoked the right of "
           "publicity and the Data Protection Act, both outside the Copyright Act and "
           "excluded by the annotation protocol; the answer is taken from the annotator."},
 "NCB-A-029": {"provision": "36(1)(g)", "answer_from": "euchay",
   "note": "Adjudicator gave s.44(1)(b), the criminal provision. The question asks "
           "whether the club is INFRINGING, which s.36(1)(g) answers. Researcher override."},
 "NCB-A-030": {"provision": "36(1)(d)+44(1)(c)", "answer_from": "eloho",
   "note": "Adjudicator gave s.44(1)(c) alone. Both cited: s.36(1)(d) establishes "
           "infringement, s.44(1)(c) that the conduct is also an offence. Eloho cited "
           "both and her answer is used. Researcher override (partial)."},
}

gold, needs_writing = [], []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    iid = r[0]
    verdict = str(r[i[VERD]] or "").strip().lower()
    prov = str(r[i[PROV]] or "").strip()
    euchay, eloho = str(r[i["Euchay_answer"]] or "").strip(), str(r[i["Eloho_answer"]] or "").strip()

    if "1st annotator" in verdict or "euchay" in verdict:
        ans, src = euchay, "Euchay"
    elif "2nd annotator" in verdict or "eloho" in verdict:
        ans, src = eloho, "Eloho"
    elif "both" in verdict and ("correct" in verdict or "right" in verdict):
        ans, src = euchay, "Euchay (both endorsed)"
    else:
        ans, src = "", "NEEDS WRITING"

    rec = {"item_id": iid, "category": "single_hop",
           "gold_provision": prov, "gold_answer": ans, "answer_source": src,
           "adjudicator_verdict": str(r[i[VERD]] or "").strip(),
           "adjudicator_reasoning": str(r[i[REAS]] or "").strip()}

    if iid in OVERRIDE:
        o = OVERRIDE[iid]
        rec["gold_provision"] = o["provision"]
        rec["gold_answer"] = eloho if o["answer_from"] == "eloho" else euchay
        rec["answer_source"] = f"{o['answer_from'].title()} (researcher override)"
        rec["override_note"] = o["note"]

    gold.append(rec)
    if rec["answer_source"] == "NEEDS WRITING":
        needs_writing.append((iid, verdict, prov))

Path("data/eval").mkdir(parents=True, exist_ok=True)
Path("data/eval/catA_gold.jsonl").write_text(
    "\n".join(json.dumps(g, ensure_ascii=False) for g in gold) + "\n", encoding="utf-8")

print(f"wrote data/eval/catA_gold.jsonl  ({len(gold)} items)")
print("answer source:", dict(Counter(g["answer_source"] for g in gold)))
print(f"\nitems needing a written gold answer: {len(needs_writing)}")
for iid, v, p in needs_writing:
    print(f"  {iid}  verdict={v!r}  provision={p}")
