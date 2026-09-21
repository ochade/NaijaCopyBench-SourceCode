import re
from pathlib import Path
from openpyxl import load_workbook

ws = load_workbook(Path.home()/"Downloads"/"catA_adjudication_sheet.xlsx").active
hdr = [c.value for c in ws[1]]
i = {h: n for n, h in enumerate(hdr)}
print(f"{'item':12} | {'Euchay':<30} | Eloho")
print("-"*80)
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[i["review_reason"]] == "PROVISION DISAGREEMENT":
        print(f"{row[i['item_id']]:12} | {str(row[i['Euchay_provision']])[:29]:<30} | {str(row[i['Eloho_provision']])[:40]}")
