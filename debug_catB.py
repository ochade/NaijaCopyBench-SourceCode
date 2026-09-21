from pathlib import Path
from openpyxl import load_workbook

DL = Path.home() / "Downloads"
for name in ["catB_annotation_sheet_Euchay.xlsx", "catB_annotation_sheet_Eloho.xlsx"]:
    wb = load_workbook(DL / name, data_only=True)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    print(f"=== {name} ===")
    print("sheet:", ws.title, "| rows:", ws.max_row)
    print("headers:", hdr)
    # first three data rows, showing item_id and provision exactly as stored
    for r in range(2, 5):
        vals = [ws.cell(row=r, column=c).value for c in range(1, len(hdr)+1)]
        print(f"  row{r}: item_id={vals[0]!r}  provision={vals[2]!r}")
    print()
