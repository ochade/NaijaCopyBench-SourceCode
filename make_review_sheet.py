"""Build a review sheet from a SINGLE annotator's file.
Usage: python make_review_sheet.py C Euchay
"""
import sys
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from collections import Counter

CAT   = sys.argv[1] if len(sys.argv) > 1 else "C"
ANNOT = sys.argv[2] if len(sys.argv) > 2 else "Euchay"
DL    = Path.home() / "Downloads"
SRC   = DL / f"cat{CAT}_annotation_sheet_{ANNOT}.xlsx"
OUT   = DL / f"cat{CAT}_review_sheet.xlsx"

NOTES = {
 "C": "DIVERGENCE. Answer from the Nigerian Act only; do not compare jurisdictions.",
 "D": "CONTROL. The provision cited matters as much as the answer.",
 "E": "VERIFICATION. Cited provisions or bodies MAY OR MAY NOT exist. Check each against the Act.",
 "F": "SUMMARISATION. List the points any correct summary must contain, and statements that would be wrong.",
}
COLMAP = {
 "C": ("provision","governing_provision","answer","answer"),
 "D": ("provision","governing_provision","answer","answer"),
 "E": ("exists?","does_the_cited_provision_or_body_exist","correct_position","correct_position"),
 "F": ("provisions","governing_provisions","required_elements","required_elements_any_correct_summary_must_contain"),
}
L5, K5, L6, K6 = COLMAP.get(CAT, ("provision","governing_provision","answer","answer"))

ws_in = load_workbook(SRC, data_only=True).active
hdr = [c.value for c in ws_in[1]]
if hdr and hdr[0] in (None, ""): hdr[0] = "item_id"
hdr = [h if h else f"col{i}" for i, h in enumerate(hdr)]

rows = []
for r in ws_in.iter_rows(min_row=2, values_only=True):
    rec = {k: (str(v).strip() if v is not None else "") for k, v in zip(hdr, r)}
    if not rec.get("item_id"): continue
    conf  = rec.get("confidence","")
    extra = rec.get("common_errors_that_would_be_wrong","") if CAT == "F" else ""
    flag  = ("LOW CONFIDENCE" if "low" in conf.lower()
             else "NOTES FLAGGED" if rec.get("notes","").strip() else "")
    rows.append([rec["item_id"], NOTES.get(CAT,""), flag, rec.get("question",""),
                 rec.get(K5,""), rec.get(K6,""), extra, conf, rec.get("notes",""),
                 "", "", "", "", ""])

order = {"LOW CONFIDENCE":0, "NOTES FLAGGED":1, "":2}
rows.sort(key=lambda r: (order.get(r[2],2), r[0]))

HDR = ["item_id","category_note","flag","question",
       f"{ANNOT}_{L5}", f"{ANNOT}_{L6}",
       (f"{ANNOT}_common_errors" if CAT == "F" else "extra"),
       f"{ANNOT}_confidence", f"{ANNOT}_notes",
       "AGREE?","FINAL_provision","FINAL_answer","reasoning","question_verdict"]

wb = Workbook(); ws = wb.active; ws.title = f"cat{CAT}_review"
ws.append(HDR)
for r in rows: ws.append(r)

for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F3864")
    c.alignment = Alignment(wrap_text=True, vertical="center")

fills = {"LOW CONFIDENCE":"FFEB9C","NOTES FLAGGED":"DDEBF7"}
for i, r in enumerate(rows, start=2):
    if r[2] in fills:
        ws.cell(row=i, column=3).fill = PatternFill("solid", fgColor=fills[r[2]])
    for col in (2,4,6,7,9,12,13):
        ws.cell(row=i, column=col).alignment = Alignment(wrap_text=True, vertical="top")

for col, w in {"A":13,"B":36,"C":16,"D":46,"E":20,"F":48,"G":38,"H":11,
               "I":32,"J":10,"K":18,"L":48,"M":36,"N":18}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "D2"
wb.save(OUT)

c = Counter(r[2] for r in rows)
print("wrote", OUT)
print(f"{len(rows)} items | annotator: {ANNOT}")
for k, v in c.items(): print(f"  {k or 'no flag'}: {v}")
