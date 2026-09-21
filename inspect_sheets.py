from openpyxl import load_workbook
from pathlib import Path

DL = Path.home() / "Downloads"
for name in ["catA_annotation_sheet_Euchay.xlsx", "catA_annotation_sheet_Eloho.xlsx"]:
    wb = load_workbook(DL / name, data_only=True)
    ws = wb.active
    print(f"=== {name} ===")
    print("sheet:", ws.title, "| rows:", ws.max_row, "| cols:", ws.max_column)
    headers = [c.value for c in ws[1]]
    print("headers:", headers)
    # show first data row so we can see what's actually filled
    row2 = [c.value for c in ws[2]]
    for h, v in zip(headers, row2):
        print(f"   {h}: {str(v)[:70] if v else '(empty)'}")
    print()
